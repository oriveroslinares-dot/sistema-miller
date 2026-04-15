import io
from flask import Blueprint, render_template, request, send_file
from flask_login import login_required
from models import Producto, InventarioActual, Kardex, Categoria
from sqlalchemy import func

reportes_bp = Blueprint('reportes', __name__, url_prefix='/reportes')


@reportes_bp.route('/stock')
@login_required
def stock():
    categoria_id = request.args.get('categoria', type=int)
    alertas = request.args.get('alertas', type=int)

    query = (InventarioActual.query
             .join(InventarioActual.producto)
             .filter(Producto.activo == True))

    if categoria_id:
        query = query.filter(Producto.categoria_id == categoria_id)

    items = query.order_by(Producto.codigo).all()

    if alertas:
        items = [i for i in items if i.stock <= i.producto.stock_minimo]

    categorias = Categoria.query.all()
    total_valor = sum(i.valor_total for i in items)

    return render_template('reportes/stock.html', items=items,
                           categorias=categorias, categoria_sel=categoria_id,
                           total_valor=total_valor, alertas=alertas)


@reportes_bp.route('/movimientos')
@login_required
def movimientos():
    desde = request.args.get('desde', '')
    hasta = request.args.get('hasta', '')
    tipo = request.args.get('tipo', '')
    categoria_id = request.args.get('categoria', type=int)

    query = Kardex.query.join(Kardex.producto)

    if desde:
        query = query.filter(Kardex.fecha >= desde)
    if hasta:
        query = query.filter(Kardex.fecha <= hasta + ' 23:59:59')
    if tipo:
        query = query.filter(Kardex.tipo == tipo)
    if categoria_id:
        query = query.filter(Producto.categoria_id == categoria_id)

    movs = query.order_by(Kardex.fecha.desc(), Kardex.id.desc()).limit(500).all()
    categorias = Categoria.query.all()

    return render_template('reportes/movimientos.html', movimientos=movs,
                           categorias=categorias, categoria_sel=categoria_id,
                           desde=desde, hasta=hasta, tipo_sel=tipo)


# ─── EXPORTAR A EXCEL ────────────────────────────────────────────────────────

@reportes_bp.route('/exportar/stock')
@login_required
def exportar_stock():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    items = (InventarioActual.query
             .join(InventarioActual.producto)
             .filter(Producto.activo == True)
             .order_by(Producto.codigo)
             .all())

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Stock Actual'

    header_fill = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    headers = ['Código', 'Producto', 'Categoría', 'Proveedor', 'Unidad',
               'Stock', 'Stock Mín.', 'Costo Prom.', 'Precio Venta', 'Valor Total']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row, i in enumerate(items, 2):
        ws.cell(row=row, column=1, value=i.producto.codigo)
        ws.cell(row=row, column=2, value=i.producto.nombre)
        ws.cell(row=row, column=3, value=i.producto.categoria.nombre)
        ws.cell(row=row, column=4, value=i.producto.proveedor or '')
        ws.cell(row=row, column=5, value=i.producto.unidad)
        ws.cell(row=row, column=6, value=i.stock)
        ws.cell(row=row, column=7, value=i.producto.stock_minimo)
        ws.cell(row=row, column=8, value=i.costo_promedio)
        ws.cell(row=row, column=9, value=i.producto.precio_venta or 0)
        ws.cell(row=row, column=10, value=i.valor_total)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='stock_miller.xlsx')


@reportes_bp.route('/exportar/movimientos')
@login_required
def exportar_movimientos():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    desde = request.args.get('desde', '')
    hasta = request.args.get('hasta', '')
    tipo = request.args.get('tipo', '')

    query = Kardex.query.join(Kardex.producto)
    if desde:
        query = query.filter(Kardex.fecha >= desde)
    if hasta:
        query = query.filter(Kardex.fecha <= hasta + ' 23:59:59')
    if tipo:
        query = query.filter(Kardex.tipo == tipo)

    movs = query.order_by(Kardex.fecha.asc(), Kardex.id.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Movimientos'

    header_fill = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    headers = ['Fecha', 'Tipo', 'Código', 'Producto', 'Proveedor', 'Empresa',
               'Documento', 'Concepto', 'Cantidad', 'Costo Unit.',
               'Costo Total', 'Precio Venta', 'Saldo Cant.', 'Costo Prom.', 'Saldo Valor', 'Usuario']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row, m in enumerate(movs, 2):
        ws.cell(row=row, column=1, value=m.fecha.strftime('%d/%m/%Y'))
        ws.cell(row=row, column=2, value=m.tipo)
        ws.cell(row=row, column=3, value=m.producto.codigo)
        ws.cell(row=row, column=4, value=m.producto.nombre)
        ws.cell(row=row, column=5, value=m.producto.proveedor or '')
        ws.cell(row=row, column=6, value=m.empresa.nombre if m.empresa else '')
        ws.cell(row=row, column=7, value=m.documento or '')
        ws.cell(row=row, column=8, value=m.concepto or '')
        ws.cell(row=row, column=9, value=m.cantidad)
        ws.cell(row=row, column=10, value=m.costo_unitario)
        ws.cell(row=row, column=11, value=m.costo_total or 0)
        ws.cell(row=row, column=12, value=m.precio_venta or '')
        ws.cell(row=row, column=13, value=m.saldo_cantidad)
        ws.cell(row=row, column=14, value=m.costo_promedio or 0)
        ws.cell(row=row, column=15, value=m.saldo_valor or 0)
        ws.cell(row=row, column=16, value=m.usuario.nombre if m.usuario else '')

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='movimientos_miller.xlsx')
